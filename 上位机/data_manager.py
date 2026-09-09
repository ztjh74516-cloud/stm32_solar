# -*- coding: utf-8 -*-
"""
数据存储与Excel导出模块 (data_manager.py)
对应 STM32 真实功能:
字段:
    - timestamp : 采集时间
    - voltage   : 输入电压 (V)
    - current   : 输入电流 (A)
    - power     : 输出功率 (W)
    - shunt     : 分流电阻采样计数值 (SH)
    - threshold : 电压保护阈值 (TH)
    - relay     : 继电器状态 (RLY: ON / OFF)
"""

import os
import threading
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


class DataManager:
    """太阳能监控数据管理器"""

    EXCEL_HEADERS = [
        "时间",
        "输入电压(V)",
        "输入电流(A)",
        "输出功率(W)",
        "分流采样值(SH)",
        "保护阈值(TH/V)",
        "继电器状态(RLY)",
    ]

    PRIMARY_COLOR = "2196F3"
    HEADER_TEXT_COLOR = "FFFFFF"
    TEXT_COLOR = "333333"
    BORDER_COLOR = "E0E0E0"

    def __init__(self, max_records: Optional[int] = 5000):
        self._records: List[Dict[str, Any]] = []
        self._lock = threading.Lock()
        self._max_records = max_records

    def add_record(self, data_dict: Dict[str, Any]) -> Dict[str, Any]:
        ts = data_dict.get("timestamp")
        if not isinstance(ts, datetime):
            ts = datetime.now()

        def _safe_float(key: str, default: float = 0.0) -> float:
            val = data_dict.get(key, default)
            try:
                return float(val)
            except (ValueError, TypeError):
                return default

        voltage = _safe_float("voltage")
        current = _safe_float("current")

        if "power" in data_dict and data_dict["power"] is not None:
            power = _safe_float("power")
        else:
            power = round(voltage * current, 4)

        shunt = data_dict.get("shunt", 0)
        try:
            shunt = int(shunt)
        except (ValueError, TypeError):
            shunt = 0

        threshold = _safe_float("threshold", 10.0)
        relay = str(data_dict.get("relay", "OFF")).upper()

        record: Dict[str, Any] = {
            "timestamp": ts,
            "voltage": voltage,
            "current": current,
            "power": power,
            "shunt": shunt,
            "threshold": threshold,
            "relay": relay,
        }

        with self._lock:
            self._records.append(record)
            if self._max_records is not None and len(self._records) > self._max_records:
                del self._records[:-self._max_records]

        return record

    def get_records(self, count: Optional[int] = None) -> List[Dict[str, Any]]:
        with self._lock:
            if count is None or count <= 0 or count >= len(self._records):
                return [dict(r) for r in self._records]
            return [dict(r) for r in self._records[-count:]]

    def get_all_records(self) -> List[Dict[str, Any]]:
        return self.get_records(None)

    def get_record_count(self) -> int:
        with self._lock:
            return len(self._records)

    def clear_records(self) -> None:
        with self._lock:
            self._records.clear()

    def get_latest(self) -> Optional[Dict[str, Any]]:
        with self._lock:
            if self._records:
                return dict(self._records[-1])
            return None

    def export_to_excel(self, filepath: str, count: Optional[int] = 100) -> bool:
        records = self.get_records(count)

        abs_path = os.path.abspath(filepath)
        dir_name = os.path.dirname(abs_path)
        if dir_name and not os.path.exists(dir_name):
            os.makedirs(dir_name, exist_ok=True)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "太阳能发电监控数据"
        ws.views.sheetView[0].showGridLines = True

        font_header = Font(name="微软雅黑", size=11, bold=True, color=self.HEADER_TEXT_COLOR)
        fill_header = PatternFill(
            start_color=self.PRIMARY_COLOR,
            end_color=self.PRIMARY_COLOR,
            fill_type="solid",
        )
        align_center = Alignment(horizontal="center", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")

        font_data = Font(name="微软雅黑", size=10, color=self.TEXT_COLOR)

        thin_border_side = Side(border_style="thin", color=self.BORDER_COLOR)
        cell_border = Border(
            top=thin_border_side,
            bottom=thin_border_side,
            left=thin_border_side,
            right=thin_border_side,
        )

        ws.row_dimensions[1].height = 28
        for col_idx, header_text in enumerate(self.EXCEL_HEADERS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header_text)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = cell_border

        for row_idx, r in enumerate(records, start=2):
            ws.row_dimensions[row_idx].height = 20

            ts_val = r.get("timestamp")
            if isinstance(ts_val, datetime):
                formatted_ts = ts_val.strftime("%Y-%m-%d %H:%M:%S")
            else:
                formatted_ts = str(ts_val or "")

            row_data = [
                (formatted_ts, align_center, None),
                (r.get("voltage", 0.0), align_right, "0.00"),
                (r.get("current", 0.0), align_right, "0.000"),
                (r.get("power", 0.0), align_right, "0.00"),
                (r.get("shunt", 0), align_right, "0"),
                (r.get("threshold", 0.0), align_right, "0.00"),
                (r.get("relay", "OFF"), align_center, None),
            ]

            for col_idx, (val, alignment, num_format) in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = font_data
                cell.alignment = alignment
                cell.border = cell_border
                if num_format:
                    cell.number_format = num_format

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                display_len = sum(2 if ord(c) > 127 else 1 for c in val_str)
                max_len = max(max_len, display_len)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        wb.save(abs_path)
        return True
